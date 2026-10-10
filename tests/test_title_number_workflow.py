from __future__ import annotations

import io
from pathlib import Path

import pytest
import streamlit as st
from streamlit.testing.v1 import AppTest

import lyrics_dashboard.converter as converter_module
import lyrics_dashboard.parser as parser_module


APP_PATH = Path(__file__).resolve().parents[1] / "app.py"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FORBIDDEN_MARKER = "forbidden-title-link"


@pytest.fixture
def workflow_records(monkeypatch):
    parsed_sources = []
    downloads: list[dict[str, object]] = []
    conversions: list[object] = []
    original_parse = parser_module.parse_lyrics
    original_download = st.download_button
    original_convert = converter_module.convert_lyrics

    def record_parse(*args, **kwargs):
        parsed = original_parse(*args, **kwargs)
        parsed_sources.append(parsed)
        return parsed

    def record_download(*args, **kwargs):
        downloads.append(dict(kwargs))
        return original_download(*args, **kwargs)

    def record_convert(*args, **kwargs):
        conversions.append(args[0])
        return original_convert(*args, **kwargs)

    monkeypatch.setattr(parser_module, "parse_lyrics", record_parse)
    monkeypatch.setattr(st, "download_button", record_download)
    monkeypatch.setattr(converter_module, "convert_lyrics", record_convert)
    return parsed_sources, downloads, conversions


def _button(app: AppTest, label: str):
    return next(button for button in app.button if button.label == label)


def _preview(app: AppTest):
    return next(area for area in app.text_area if area.label == "Converted lyrics")


def _docx_source(source: str) -> bytes:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    paragraph = document.add_paragraph()
    relationship_id = paragraph.part.relate_to(
        f"https://hidden.invalid/{FORBIDDEN_MARKER}-destination",
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = f"Open song website {FORBIDDEN_MARKER}-display"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    for line in source.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _submit_source(input_kind: str, source: str) -> AppTest:
    app = AppTest.from_file(str(APP_PATH), default_timeout=20).run()
    if input_kind == "Paste text":
        next(radio for radio in app.radio if radio.label == "Input source").set_value(
            "Paste text"
        ).run(timeout=20)
        next(area for area in app.text_area if area.label == "Paste song text").set_value(
            source
        )
        _button(app, "PROCESS TEXT").click().run(timeout=20)
    else:
        app.file_uploader[0].set_value(
            ("unrelated source filename.docx", _docx_source(source), DOCX_MIME)
        ).run(timeout=20)
    assert not app.exception
    assert not app.error
    return app


@pytest.mark.parametrize("input_kind", ["Upload DOCX", "Paste text"])
@pytest.mark.parametrize(
    ("source_title", "expected_title", "language"),
    [
        ("135 人算什麼 (分2)", "135 人算什麼 分2", "zh"),
        ("001 Groot is Uw trouw", "001 Groot is Uw trouw", "nl"),
        ("42 Holy Forever", "42 Holy Forever", "nl"),
        (
            "135 人算什麼 (分2) | What Is Man (Part 2)",
            "135 人算什麼 分2 What Is Man Part 2",
            "bilingual",
        ),
        (
            "001 What Is Man (Part 2) | 人算什麼 (分2)",
            "001 人算什麼 分2 What Is Man Part 2",
            "bilingual",
        ),
        ("Holy Forever", "Holy Forever", "nl"),
    ],
    ids=[
        "chinese",
        "dutch-leading-zeros",
        "english",
        "bilingual",
        "latin-first-bilingual-leading-zeros",
        "unnumbered",
    ],
)
def test_title_number_survives_input_preview_exact_update_and_utf8_download(
    workflow_records,
    input_kind: str,
    source_title: str,
    expected_title: str,
    language: str,
) -> None:
    parsed_sources, downloads, conversions = workflow_records
    if language == "zh":
        lyrics = "Verse 1\n賜我平安\n\nChorus 1\n感謝天父\n"
    elif language == "bilingual":
        lyrics = (
            "Verse 1\n賜我平安\n\nChorus 1\n感謝天父\n\n"
            "Verse 2\nGive me peace\n\nChorus 2\nThank You Father\n"
        )
    else:
        lyrics = "Verse 1\nA short song\n\nChorus 1\nWe sing together\n"
    source = (
        f"\nhttps://example.invalid/{FORBIDDEN_MARKER}\n\n"
        f"{source_title}\n\n{lyrics}"
    )
    app = _submit_source(input_kind, source)
    assert parsed_sources[-1].raw_title == source_title
    assert FORBIDDEN_MARKER not in repr(parsed_sources[-1])
    if language == "bilingual":
        _button(app, "Validate these manual matches").click().run(timeout=20)
    assert not app.exception
    assert not app.error

    generated = _preview(app).value
    assert generated.splitlines()[:2] == ["[Title]", expected_title]
    assert "[Verse 1]" in generated
    assert "[Chorus 1]" in generated
    assert generated.count("[Title]") == 1
    assert FORBIDDEN_MARKER not in generated
    assert app.session_state["generated_output"] == generated
    assert app.session_state["editable_preview"] == generated
    assert app.session_state["saved_final_output"] == generated
    assert downloads[-1]["data"] == generated.encode("utf-8")
    assert downloads[-1]["data"].decode("utf-8") == generated
    assert downloads[-1]["file_name"] == f"{expected_title}.txt"
    assert downloads[-1]["mime"] == "text/plain; charset=utf-8"
    assert app.get("download_button")[0].proto.ignore_rerun is True

    # Edit title text while retaining its number and deliberately unnormalized text.
    edited_title = f"{expected_title}  Eigen  titel!  "
    exact_edit = generated.replace(
        f"[Title]\n{expected_title}\n",
        f"[Title]\n{edited_title}\n",
        1,
    ) + "\n  Handmatige  tekst?! // 自訂  \n\n"
    conversion_count = len(conversions)
    _preview(app).set_value(exact_edit).run(timeout=20)
    assert not app.exception
    assert _preview(app).value == exact_edit
    assert app.session_state["saved_final_output"] == generated
    assert downloads[-1]["data"] == generated.encode("utf-8")
    assert downloads[-1]["file_name"] == f"{expected_title}.txt"

    _button(app, "UPDATE").click().run(timeout=20)
    assert not app.exception
    assert _preview(app).value == exact_edit
    assert _preview(app).value.splitlines()[1] == edited_title
    assert app.session_state["generated_output"] == generated
    assert app.session_state["editable_preview"] == exact_edit
    assert app.session_state["saved_final_output"] == exact_edit
    assert downloads[-1]["data"] == exact_edit.encode("utf-8")
    assert downloads[-1]["data"].decode("utf-8") == exact_edit
    assert downloads[-1]["file_name"] == f"{expected_title} Eigen titel!.txt"
    assert len(conversions) == conversion_count
